"""
Script 1: extract_to_staging_db.py
Système d'Extraction Intelligent : Rapports Non Structurés -> Tables de Staging SQL
Version Groq API
"""

import os
import re
import json
import mysql.connector
from pathlib import Path
from typing import Dict, List, Optional
from groq import Groq
from dotenv import load_dotenv
import uuid

# Imports pour la lecture de différents formats de fichiers
PDF_LIB = None
PDF_AVAILABLE = False
try:
    import PyPDF2
    PDF_LIB = PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    try:
        import pypdf
        PDF_LIB = pypdf
        PDF_AVAILABLE = True
    except ImportError:
        PDF_AVAILABLE = False
        print("⚠️  Attention: PyPDF2/pypdf non installé. Support PDF désactivé.")

DOCX_AVAILABLE = False
DOCX_Document = None
try:
    from docx import Document
    DOCX_Document = Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("⚠️  Attention: python-docx non installé. Support DOCX désactivé.")

# --- Configuration Globale ---
load_dotenv()
GROQ_API_KEY = os.getenv("API_KEY")  # Utilise API_KEY depuis .env
DB_CONFIG = {
    'host': os.getenv("DB_HOST"), 
    'user': os.getenv("DB_USER"),
    'password': os.getenv("DB_PASS"), 
    'database': os.getenv("DB_NAME"),
    'charset': 'utf8mb4'
}
# Utiliser les rapports générés par structure_to_pdf_text.py
# Support multiple formats: txt, pdf, docx
REPORTS_FOLDER = Path("moroccan_unstructured_data")
SUPPORTED_EXTENSIONS = ['.txt', '.pdf', '.docx', '.doc']


def extract_text_from_pdf(file_path: Path) -> Optional[str]:
    """Extrait le texte d'un fichier PDF"""
    if not PDF_AVAILABLE or PDF_LIB is None:
        print(f"  ⚠️  Support PDF non disponible. Installez PyPDF2: pip install PyPDF2")
        return None
    
    try:
        text_parts = []
        with open(file_path, 'rb') as file:
            pdf_reader = PDF_LIB.PdfReader(file)
            for page in pdf_reader.pages:
                text_parts.append(page.extract_text())
        
        return '\n'.join(text_parts)
    except Exception as e:
        print(f"  ❌ Erreur lors de l'extraction PDF: {e}")
        return None


def extract_text_from_docx(file_path: Path) -> Optional[str]:
    """Extrait le texte d'un fichier DOCX"""
    if not DOCX_AVAILABLE or DOCX_Document is None:
        print(f"  ⚠️  Support DOCX non disponible. Installez python-docx: pip install python-docx")
        return None
    
    try:
        doc = DOCX_Document(file_path)
        text_parts = []
        for paragraph in doc.paragraphs:
            text_parts.append(paragraph.text)
        return '\n'.join(text_parts)
    except Exception as e:
        print(f"  ❌ Erreur lors de l'extraction DOCX: {e}")
        return None


def extract_text_from_file(file_path: Path) -> Optional[str]:
    """Extrait le texte d'un fichier selon son extension"""
    extension = file_path.suffix.lower()
    
    if extension == '.txt':
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print(f"  ❌ Erreur lors de la lecture TXT: {e}")
            return None
    elif extension == '.pdf':
        return extract_text_from_pdf(file_path)
    elif extension in ['.docx', '.doc']:
        return extract_text_from_docx(file_path)
    else:
        print(f"  ⚠️  Format non supporté: {extension}")
        return None


def extract_patient_metadata_from_text(report_text: str) -> Dict:
    marker = "METADONNEES_PATIENT_JSON:"
    idx = report_text.find(marker)
    if idx == -1:
        return {}
    start = idx + len(marker)
    length = len(report_text)
    while start < length and report_text[start] in " \n\r\t":
        start += 1
    if start >= length or report_text[start] != '{':
        return {}
    depth = 0
    end = start
    for i in range(start, length):
        ch = report_text[i]
        if ch == '{':
            depth += 1
        elif ch == '}':
            depth -= 1
            if depth == 0:
                end = i + 1
                break
    if depth != 0:
        return {}
    json_block = report_text[start:end]
    try:
        # Nettoyer NaN, Infinity, etc. avant le parsing
        json_block = re.sub(r'\bNaN\b', 'null', json_block)
        json_block = re.sub(r'\bInfinity\b', 'null', json_block)
        json_block = re.sub(r'\b-infinity\b', 'null', json_block, flags=re.IGNORECASE)
        metadata = json.loads(json_block)
        if isinstance(metadata, dict):
            return metadata
    except json.JSONDecodeError:
        return {}
    return {}


class StagingExtractor:
    """Extrait les données et les charge dans les tables de staging"""
    
    def __init__(self, api_key: str, reports_folder: Path, db_config: dict):
        self.reports_folder = reports_folder
        
        # Configuration de Groq
        self.client = Groq(api_key=api_key)
        # Utilise llama-3.1-8b-instant : rapide, performant et économique
        self.model_name = "llama-3.1-8b-instant"
        
        try:
            self.db_conn = mysql.connector.connect(**db_config)
            self.db_cursor = self.db_conn.cursor(dictionary=True)
            print("\n" + "="*70)
            print("✅ CONNEXION À LA BASE DE DONNÉES RÉUSSIE")
            print("="*70)
            
            # VÉRIFICATION: Tester la connexion avec une requête simple
            self.db_cursor.execute("SELECT DATABASE() as db_name, CONNECTION_ID() as conn_id, USER() as user")
            conn_info = self.db_cursor.fetchone()
            print(f"📊 Base de données: {conn_info.get('db_name', 'N/A')}")
            print(f"🔌 ID Connexion: {conn_info.get('conn_id', 'N/A')}")
            print(f"👤 Utilisateur: {conn_info.get('user', 'N/A')}")
            
            # Vérifier que les tables existent
            print("\n🔍 Recherche des tables de staging...")
            self.db_cursor.execute("SHOW TABLES LIKE 'staging_%'")
            tables = [row[list(row.keys())[0]] for row in self.db_cursor.fetchall()]
            
            if tables:
                print(f"✅ {len(tables)} table(s) de staging trouvée(s):")
                for i, table in enumerate(tables, 1):
                    # Vérifier le nombre de lignes dans chaque table
                    try:
                        self.db_cursor.execute(f"SELECT COUNT(*) as count FROM {table}")
                        count = self.db_cursor.fetchone()['count']
                        print(f"   {i}. {table} ({count} ligne(s))")
                    except:
                        print(f"   {i}. {table}")
                print("="*70 + "\n")
            else:
                print("⚠️  ATTENTION: Aucune table staging_* trouvée!")
                print("   Vérifiez que le schéma SQL a été exécuté correctement.")
                print("="*70 + "\n")
                
        except mysql.connector.Error as err:
            print("\n" + "="*70)
            print("❌ ERREUR DE CONNEXION À LA BASE DE DONNÉES")
            print("="*70)
            print(f"Erreur: {err}")
            print("Vérifiez vos paramètres de connexion dans le fichier .env")
            print("="*70 + "\n")
            exit()
            
        print(f"✓ Dossier des rapports : {self.reports_folder}")
        print(f"✓ Formats supportés : {', '.join(SUPPORTED_EXTENSIONS)}")
    
    def extract_with_llm(self, report_text: str, extraction_type: str) -> Dict:
        """Utilise Groq AI pour extraire les données structurées du rapport"""
        
        # Prompts optimisés pour Groq
#         prompts = {
#             'patient_info': """Extrait les informations du patient du rapport médical. 
# Retourne UNIQUEMENT un objet JSON valide avec ces champs:
# {"id", "first_name", "last_name", "birthdate": "YYYY-MM-DD", "gender": "M ou F", "ssn", "address", "city", "state", "zip", "race", "ethnicity"}
# Ne retourne que le JSON, sans texte supplémentaire.""",
            
#             'encounter': """Extrait les informations de la consultation du rapport médical.
# Retourne UNIQUEMENT un objet JSON valide avec ces champs:
# {"id", "start_datetime": "YYYY-MM-DD HH:MM:SS", "stop_datetime", "patient_id", "organization_id", "provider_id", "encounter_class", "code", "description", "base_encounter_cost", "total_claim_cost", "payer_coverage", "reason_description"}
# Ne retourne que le JSON, sans texte supplémentaire.""",
            
#             'conditions': """Extrait TOUTES les pathologies/diagnostics du rapport médical.
# Retourne UNIQUEMENT un array JSON valide:
# [{"patient_id", "encounter_id", "start_date": "YYYY-MM-DD", "stop_date", "code", "description"}, ...]
# Ne retourne que le JSON, sans texte supplémentaire.""",
            
#             'medications': """Extrait TOUTES les prescriptions médicamenteuses du rapport médical.
# Retourne UNIQUEMENT un array JSON valide:
# [{"patient_id", "encounter_id", "start_datetime", "stop_datetime", "code", "description", "base_cost", "total_cost", "payer_coverage", "reason_description"}, ...]
# Ne retourne que le JSON, sans texte supplémentaire.""",
            
#             'observations': """Extrait TOUTES les observations/mesures vitales du rapport médical.
# Retourne UNIQUEMENT un array JSON valide:
# [{"patient_id", "encounter_id", "date_recorded", "code", "description", "value", "units", "type"}, ...]
# Ne retourne que le JSON, sans texte supplémentaire.""",
            
#             'allergies': """Extrait TOUTES les allergies du rapport médical.
# Retourne UNIQUEMENT un array JSON valide:
# [{"patient_id", "encounter_id", "start_date", "stop_date", "code", "description"}, ...]
# Ne retourne que le JSON, sans texte supplémentaire.""",
            
#             'procedures': """Extrait TOUS les actes médicaux/interventions du rapport médical.
# Retourne UNIQUEMENT un array JSON valide:
# [{"patient_id", "encounter_id", "date_performed", "code", "description", "base_cost", "reason_description"}, ...]
# Ne retourne que le JSON, sans texte supplémentaire.""",
            
#             'immunizations': """Extrait TOUTES les vaccinations du rapport médical.
# Retourne UNIQUEMENT un array JSON valide:
# [{"patient_id", "encounter_id", "date_administered", "code", "description", "base_cost"}, ...]
# Ne retourne que le JSON, sans texte supplémentaire."""
#         }

        prompts = {
    'patient_info': """Extrait les informations du patient du rapport médical. 
Retourne UNIQUEMENT un objet JSON valide avec ces champs:
{"id", "first_name", "last_name", "birthdate": "YYYY-MM-DD", "gender": "M ou F", "ssn", "address", "city", "state", "zip", "race", "ethnicity"}
Ne retourne que le JSON, sans texte supplémentaire.""",
    
    'encounter': """Extrait les informations de la consultation du rapport médical.
Retourne UNIQUEMENT un objet JSON valide avec ces champs:
{"id", "start_datetime": "YYYY-MM-DD HH:MM:SS", "stop_datetime": "YYYY-MM-DD HH:MM:SS", "patient_id", "organization_id", "provider_id", "encounter_class", "code", "description", "base_encounter_cost", "total_claim_cost", "payer_coverage", "reason_description"}
Ne retourne que le JSON, sans texte supplémentaire.""",
    
    'conditions': """Extrait TOUTES les pathologies/diagnostics du rapport médical.
Retourne UNIQUEMENT un array JSON valide:
[{"start_date": "YYYY-MM-DD", "stop_date": "YYYY-MM-DD", "code", "description"}, ...]
Ne retourne que le JSON, sans texte supplémentaire. NE PAS INCLURE patient_id ou encounter_id.""",
    
    'medications': """Extrait TOUTES les prescriptions médicamenteuses du rapport médical.
Retourne UNIQUEMENT un array JSON valide:
[{"start_datetime": "YYYY-MM-DD HH:MM:SS", "stop_datetime": "YYYY-MM-DD HH:MM:SS", "code", "description", "base_cost", "total_cost", "payer_coverage", "reason_description"}, ...]
Ne retourne que le JSON, sans texte supplémentaire. NE PAS INCLURE patient_id ou encounter_id.""",
    
    'observations': """Extrait TOUTES les observations/mesures vitales du rapport médical.
Retourne UNIQUEMENT un array JSON valide:
[{"date_recorded": "YYYY-MM-DD HH:MM:SS", "code", "description", "value", "units", "type"}, ...]
Ne retourne que le JSON, sans texte supplémentaire. NE PAS INCLURE patient_id ou encounter_id.""",
    
    'allergies': """Extrait TOUTES les allergies du rapport médical.
Retourne UNIQUEMENT un array JSON valide:
[{"start_date": "YYYY-MM-DD", "stop_date": "YYYY-MM-DD", "code", "description"}, ...]
Ne retourne que le JSON, sans texte supplémentaire. NE PAS INCLURE patient_id ou encounter_id.""",
    
    'procedures': """Extrait TOUS les actes médicaux/interventions du rapport médical.
Retourne UNIQUEMENT un array JSON valide:
[{"date_performed": "YYYY-MM-DD HH:MM:SS", "code", "description", "base_cost", "reason_description"}, ...]
Ne retourne que le JSON, sans texte supplémentaire. NE PAS INCLURE patient_id ou encounter_id.""",
    
    'immunizations': """Extrait TOUTES les vaccinations du rapport médical.
Retourne UNIQUEMENT un array JSON valide:
[{"date_administered": "YYYY-MM-DD HH:MM:SS", "code", "description", "base_cost"}, ...]
Ne retourne que le JSON, sans texte supplémentaire. NE PAS INCLURE patient_id ou encounter_id."""
}
        
        prompt = prompts.get(extraction_type)
        if not prompt: 
            return {}

        try:
            # Appel à Groq
            full_prompt = f"{prompt}\n\nRAPPORT MÉDICAL:\n{report_text[:12000]}"
            response = self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": "Tu es un assistant expert en extraction de données médicales. Tu retournes uniquement du JSON valide."},
                    {"role": "user", "content": full_prompt}
                ],
                temperature=0.1,  # Basse température pour plus de cohérence
                max_tokens=4000
            )
            
            # Extraction du JSON
            content = response.choices[0].message.content
            
            # Nettoyage du markdown si présent
            content = re.sub(r'```json\s*', '', content)
            content = re.sub(r'```\s*', '', content)
            content = content.strip()
            
            # Extraction du JSON
            json_match = re.search(r'\{.*\}|\[.*\]', content, re.DOTALL)
            if json_match:
                json_str = json_match.group()
                # Nettoyer NaN, Infinity, etc. avant le parsing
                json_str = re.sub(r'\bNaN\b', 'null', json_str)
                json_str = re.sub(r'\bInfinity\b', 'null', json_str)
                json_str = re.sub(r'\b-infinity\b', 'null', json_str, flags=re.IGNORECASE)
                try:
                    return json.loads(json_str)
                except json.JSONDecodeError as e:
                    print(f"  ⚠️  Erreur parsing JSON pour {extraction_type}: {e}")
                    return {} if extraction_type in ['patient_info', 'encounter'] else []
            
            print(f"  ⚠️  Aucun JSON valide retourné par Groq pour {extraction_type}")
            return {}
                
        except Exception as e:
            print(f"✗ Erreur extraction Groq ({extraction_type}): {e}")
            return {} if extraction_type in ['patient_info', 'encounter'] else []
    
    def clean_value(self, value):
        """Nettoie une valeur pour l'insertion en base de données"""
        import math
        
        # Si None, retourner None
        if value is None:
            return None
        
        # Si c'est un float NaN
        if isinstance(value, float) and math.isnan(value):
            return None
        
        # Si c'est une chaîne représentant NaN
        if isinstance(value, str):
            value_lower = value.lower().strip()
            if value_lower in ['nan', 'none', 'null', '', 'undefined']:
                return None
        
        # Si c'est un nombre infini
        if isinstance(value, float) and (math.isinf(value) or math.isnan(value)):
            return None
        
        return value
    
    def validate_data(self, data: Dict, table_name: str) -> Dict:
        """Valide et nettoie les données avant insertion selon le schéma exact"""
        if not isinstance(data, dict):
            print(f"  [DEBUG] validate_data: data n'est pas un dictionnaire")
            return {}
        
        # Récupérer les colonnes valides de la table
        valid_cols = self.get_table_columns(table_name)
        
        if not valid_cols:
            print(f"  ❌ ERREUR: Impossible de récupérer les colonnes de {table_name}")
            return {}
        
        # Filtrer et nettoyer les valeurs selon les colonnes réelles
        item_filtered = {}
        skipped_cols = []
        
        for k, v in data.items():
            # S'assurer que la clé est valide et n'est pas NaN
            if k in valid_cols and not (isinstance(k, str) and k.lower() == 'nan'):
                # Nettoyer la valeur
                cleaned_value = self.clean_value(v)
                # Ne pas inclure les clés avec des valeurs invalides comme nom de colonne
                item_filtered[k] = cleaned_value
            else:
                if k not in valid_cols:
                    skipped_cols.append(f"{k} (colonne inexistante)")
                elif isinstance(k, str) and k.lower() == 'nan':
                    skipped_cols.append(f"{k} (clé NaN)")
        
        if skipped_cols and table_name == 'staging_patients':
            print(f"  [DEBUG] Colonnes ignorées: {', '.join(skipped_cols[:5])}")
        
        if not item_filtered:
            print(f"  ❌ ERREUR: Aucune colonne valide après filtrage pour {table_name}")
            print(f"     [DEBUG] Colonnes reçues: {list(data.keys())}")
            print(f"     [DEBUG] Colonnes valides dans la table: {sorted(list(valid_cols))[:10]}")
        
        return item_filtered
    
    def get_primary_key(self, table_name: str) -> str:
        """Détecte automatiquement la clé primaire d'une table selon le schéma"""
        # Selon le schéma SQL :
        # - staging_patients et staging_encounters : utilisent 'id' (TEXT avec UNIQUE KEY)
        # - Toutes les autres tables : utilisent 'staging_id' (INT AUTO_INCREMENT PRIMARY KEY)
        if table_name in ['staging_patients', 'staging_encounters']:
            return 'id'
        else:
            return 'staging_id'
    
    def get_table_columns(self, table_name: str) -> set:
        """Récupère la liste des colonnes valides d'une table"""
        try:
            self.db_cursor.execute(f"DESCRIBE {table_name}")
            return {row['Field'] for row in self.db_cursor.fetchall()}
        except Exception as e:
            print(f"  ⚠️  Erreur lors de la récupération des colonnes de {table_name}: {e}")
            return set()
    
    def verify_insertion(self, table_name: str, primary_key: str, primary_value: str) -> bool:
        """Vérifie qu'un enregistrement a bien été inséré en faisant un SELECT"""
        try:
            # Construire la requête SELECT
            select_sql = f"SELECT * FROM {table_name} WHERE {primary_key} = %s LIMIT 1"
            self.db_cursor.execute(select_sql, (primary_value,))
            result = self.db_cursor.fetchone()
            
            if result:
                # Afficher quelques informations clés pour confirmation
                if table_name == 'staging_patients':
                    print(f"      ✓ Vérifié: Patient {result.get('first_name', '')} {result.get('last_name', '')} (ID: {primary_value})")
                elif table_name == 'staging_encounters':
                    print(f"      ✓ Vérifié: Consultation {result.get('id', 'N/A')} pour patient {result.get('patient_id', 'N/A')}")
                else:
                    # Pour les tables enfants, afficher le staging_id et quelques infos
                    staging_id = result.get('staging_id', primary_value)
                    desc = result.get('description', '')[:50] if result.get('description') else 'N/A'
                    print(f"      ✓ Vérifié: {table_name} - staging_id: {staging_id} ({desc}...)")
                return True
            else:
                print(f"      ⚠️  ATTENTION: L'enregistrement n'a pas été trouvé après insertion!")
                return False
        except Exception as e:
            print(f"      ⚠️  Erreur lors de la vérification: {e}")
            return False
    
    def insert_single_record(self, table_name: str, data: Dict, report_filename: str) -> bool:
        """Insère UNE seule donnée avec validation, commit immédiat et vérification"""
        if not data:
            return False
        
        # Ajouter le nom du fichier source
        data['report_filename'] = report_filename
        
        # Valider et nettoyer les données
        item_filtered = self.validate_data(data, table_name)
        
        if not item_filtered:
            print(f"  ❌ ERREUR: Aucune donnée valide pour {table_name}")
            print(f"     [DEBUG] Données reçues: {list(data.keys()) if data else 'AUCUNE'}")
            print(f"     [DEBUG] Colonnes filtrées: {len(item_filtered)} colonne(s)")
            return False
        
        cols = ', '.join(item_filtered.keys())
        placeholders = ', '.join(['%s'] * len(item_filtered))
        sql = f"INSERT INTO {table_name} ({cols}) VALUES ({placeholders})"
        
        # Détecter la clé primaire de la table
        primary_key = self.get_primary_key(table_name)
        
        try:
            # Afficher le nom de la table avant insertion
            print(f"      📋 TABLE: {table_name}")
            
            # Compter les lignes AVANT insertion
            self.db_cursor.execute(f"SELECT COUNT(*) as count_before FROM {table_name}")
            count_before = self.db_cursor.fetchone()['count_before']
            
            # INSERT
            self.db_cursor.execute(sql, list(item_filtered.values()))
            rows_affected = self.db_cursor.rowcount
            
            # VÉRIFICATION: Récupérer la clé primaire AVANT le commit
            primary_value = None
            
            if primary_key == 'staging_id':
                # Pour les tables avec staging_id AUTO_INCREMENT, récupérer le dernier ID inséré
                self.db_cursor.execute("SELECT LAST_INSERT_ID() as last_id")
                result = self.db_cursor.fetchone()
                if result and result['last_id']:
                    primary_value = result['last_id']
            else:
                # Pour staging_patients et staging_encounters, utiliser l'ID du dictionnaire
                primary_value = item_filtered.get('id')
            
            # Commit après avoir récupéré l'ID
            self.db_conn.commit()  # Commit immédiat après chaque insertion
            
            # Compter les lignes APRÈS insertion
            self.db_cursor.execute(f"SELECT COUNT(*) as count_after FROM {table_name}")
            count_after = self.db_cursor.fetchone()['count_after']
            rows_inserted = count_after - count_before
            
            # AFFICHAGE DE LA RÉUSSITE avec nom de table
            if rows_inserted > 0:
                print(f"      ✅ INSERTION RÉUSSIE dans la table: {table_name}")
                print(f"         → Lignes avant: {count_before} | Après: {count_after} | Ajoutées: +{rows_inserted}")
                print(f"         → Lignes affectées par INSERT: {rows_affected}")
                if primary_value:
                    print(f"         → Clé primaire: {primary_key} = {primary_value}")
            else:
                print(f"      ⚠️  INSERTION: Aucune nouvelle ligne ajoutée dans la table: {table_name}")
                print(f"         → Lignes avant: {count_before} | Après: {count_after}")
            
            # Vérification avec SELECT pour confirmer
            if primary_value:
                verified = self.verify_insertion(table_name, primary_key, str(primary_value))
                if verified:
                    print(f"      ✅ VÉRIFICATION: Enregistrement confirmé dans la base de données")
            else:
                # Si pas d'ID, vérifier avec le nom du fichier
                if primary_key == 'staging_id':
                    select_sql = f"SELECT * FROM {table_name} WHERE report_filename = %s ORDER BY staging_id DESC LIMIT 1"
                else:
                    select_sql = f"SELECT * FROM {table_name} WHERE report_filename = %s ORDER BY {primary_key} DESC LIMIT 1"
                
                self.db_cursor.execute(select_sql, (report_filename,))
                result = self.db_cursor.fetchone()
                if result:
                    pk_val = result.get(primary_key, 'N/A')
                    print(f"      ✅ VÉRIFICATION: Dernier enregistrement trouvé ({primary_key}: {pk_val})")
                else:
                    print(f"      ⚠️  VÉRIFICATION: Aucun enregistrement trouvé pour {report_filename}")
            
            return True
        except mysql.connector.Error as err:
            self.db_conn.rollback()  # Rollback en cas d'erreur
            if err.errno == 1062:  # Clé dupliquée
                print(f"  ⚠️  Doublon ignoré pour {table_name}: {data.get('id', 'N/A')}")
                print(f"     [INFO] L'enregistrement existe déjà dans la base de données")
            else:
                print(f"  ❌ ERREUR insertion dans {table_name}: {err}")
                print(f"     [DEBUG] Code erreur: {err.errno}")
                print(f"     [DEBUG] Message: {err.msg}")
                print(f"     [DEBUG] Colonnes tentées: {list(item_filtered.keys())}")
                print(f"     [DEBUG] Nombre de valeurs: {len(item_filtered)}")
                if item_filtered:
                    print(f"     [DEBUG] Exemple de valeur: {list(item_filtered.items())[:3]}")
            return False
    
    def insert_staging_data(self, table_name: str, data: Dict | List[Dict], report_filename: str) -> int:
        """Insère les données extraites une par une avec validation"""
        if not data: 
            return 0
        data_list = data if isinstance(data, list) else [data]
        
        count = 0
        for idx, item in enumerate(data_list, 1):
            if not isinstance(item, dict): 
                print(f"  ⚠️  Item {idx} ignoré (n'est pas un dictionnaire)")
                continue
            
            if self.insert_single_record(table_name, item, report_filename):
                count += 1
                if len(data_list) > 1:
                    print(f"    ✓ {table_name} #{idx}/{len(data_list)} inséré")
        
        return count

    def process_report(self, report_path: Path) -> Dict[str, int]:
        """Traite un rapport complet : extraction → validation → insertion immédiate"""
        print(f"\n📄 Traitement : {report_path.name} ({report_path.suffix})")
        
        # Extraire le texte selon le format du fichier
        report_text = extract_text_from_file(report_path)
        if not report_text:
            print(f"  ✗ Impossible d'extraire le texte du fichier")
            return {}
        
        if len(report_text.strip()) == 0:
            print(f"  ⚠️  Fichier vide ou aucun texte extrait")
            return {}
        
        stats = {}

        # ========== ÉTAPE 1: PATIENT ==========
        print("  🔍 [1/8] Extraction des données patient...")
        patient_metadata = extract_patient_metadata_from_text(report_text)
        patient_data = patient_metadata.copy() if patient_metadata else {}

        # Si pas de métadonnées valides, utiliser l'IA
        if not patient_data:
            print("    → Utilisation de Groq pour extraire les données patient...")
            patient_data = self.extract_with_llm(report_text, 'patient_info') or {}

        # Toujours générer un ID unique si manquant
        if not patient_data.get('id'):
            patient_data['id'] = str(uuid.uuid4())

        # Assurer la cohérence minimale
        for key in ["first_name", "last_name", "birthdate", "gender"]:
            patient_data.setdefault(key, None)

        # Validation et insertion immédiate du patient
        if patient_data:
            print(f"    → Insertion dans la table: staging_patients")
            print(f"    [DEBUG] Données patient avant insertion: {list(patient_data.keys())}")
            print(f"    [DEBUG] ID patient: {patient_data.get('id', 'MANQUANT')}")
            print(f"    [DEBUG] Nom: {patient_data.get('first_name', 'N/A')} {patient_data.get('last_name', 'N/A')}")
            
            stats['patients'] = self.insert_staging_data(
                'staging_patients', patient_data, report_path.name
            )
            
            if stats['patients'] > 0:
                print(f"    ✅ Patient inséré avec succès (ID: {patient_data['id']})")
            else:
                print("    ❌ ÉCHEC insertion patient - Arrêt du traitement")
                print("    [DEBUG] Vérifiez les logs ci-dessus pour voir la cause de l'échec")
                return stats
        else:
            print("    ❌ Aucune donnée patient extraite - Arrêt du traitement")
            print("    [DEBUG] Vérifiez que le rapport contient des métadonnées ou que Groq a extrait les données")
            return stats

        # ========== ÉTAPE 2: ENCOUNTER ==========
        print("  🔍 [2/8] Extraction des données consultation...")
        encounter_data = self.extract_with_llm(report_text, 'encounter') or {}

        if not encounter_data.get('id'):
            encounter_data['id'] = str(uuid.uuid4())
        encounter_data['patient_id'] = patient_data['id']

        # Compléter valeurs minimales
        encounter_data.setdefault("encounter_class", "wellness")
        encounter_data.setdefault("description", "Consultation de routine")
        encounter_data.setdefault("base_encounter_cost", 591)
        encounter_data.setdefault("total_claim_cost", 591)

        # Validation et insertion immédiate de l'encounter
        if encounter_data:
            print(f"    → Insertion dans la table: staging_encounters")
            stats['encounters'] = self.insert_staging_data(
                'staging_encounters', encounter_data, report_path.name
            )
            if stats['encounters'] > 0:
                print(f"    ✓ Consultation insérée (ID: {encounter_data['id']})")
            else:
                print("    ⚠️  Échec insertion consultation")
        else:
            print("    ⚠️  Aucune donnée consultation extraite")

        # ========== ÉTAPES 3-8: TABLES ENFANTS (une par une) ==========
        child_tables = [
            ('conditions', 'staging_conditions'),
            ('medications', 'staging_medications'),
            ('observations', 'staging_observations'),
            ('allergies', 'staging_allergies'),
            ('procedures', 'staging_procedures'),
            ('immunizations', 'staging_immunizations')
        ]

        for idx, (extraction_type, table_name) in enumerate(child_tables, 3):
            print(f"  🔍 [{idx}/8] Extraction {extraction_type}...")
            child_data_list = self.extract_with_llm(report_text, extraction_type)
            
            if not child_data_list or not isinstance(child_data_list, list):
                print(f"    ℹ️  Aucune donnée {extraction_type} trouvée")
                stats[extraction_type] = 0
                continue
            
            # Ajouter les IDs de référence à chaque item
            for item in child_data_list:
                if isinstance(item, dict):
                    item['patient_id'] = patient_data['id']
                    item['encounter_id'] = encounter_data.get('id')
            
            # Insertion une par une avec validation
            print(f"    → Insertion dans la table: {table_name} ({len(child_data_list)} enregistrement(s))")
            stats[extraction_type] = self.insert_staging_data(
                table_name, child_data_list, report_path.name
            )
            if stats[extraction_type] > 0:
                print(f"    ✓ {stats[extraction_type]}/{len(child_data_list)} {extraction_type} inséré(s)")

        print(f"  ✅ Résumé: {stats}")
        return stats

    
    def process_all_reports(self, max_reports: int = None):
        """Traite tous les rapports du dossier (txt, pdf, docx)"""
        # Recherche récursive de tous les fichiers supportés
        report_files = []
        for ext in SUPPORTED_EXTENSIONS:
            found_files = list(self.reports_folder.rglob(f"*{ext}"))
            report_files.extend(found_files)
        
        if not report_files:
            print(f"✗ Aucun fichier supporté trouvé dans : {self.reports_folder.resolve()}")
            print(f"   Formats recherchés: {', '.join(SUPPORTED_EXTENSIONS)}")
            return
        
        # Trier par nom pour un traitement ordonné
        report_files.sort(key=lambda x: x.name)
        
        if max_reports: 
            report_files = report_files[:max_reports]
        
        # Statistiques par format
        format_stats = {}
        for file in report_files:
            ext = file.suffix.lower()
            format_stats[ext] = format_stats.get(ext, 0) + 1
        
        print(f"\n{'='*70}\nTRAITEMENT DE {len(report_files)} RAPPORTS VERS STAGING\n{'='*70}")
        print(f"📊 Répartition par format:")
        for ext, count in sorted(format_stats.items()):
            print(f"   {ext}: {count} fichier(s)")
        print(f"{'='*70}")
        
        for idx, report_path in enumerate(report_files, 1):
            print(f"\n[{idx}/{len(report_files)}]")
            try:
                self.process_report(report_path)
            except Exception as e:
                print(f"✗ Erreur traitement {report_path.name}: {e}")
        print(f"\n{'='*70}\nEXTRACTION VERS STAGING TERMINÉE\n{'='*70}\n")
    
    def close(self):
        self.db_cursor.close()
        self.db_conn.close()
        print("✓ Connexion DB (Staging) fermée")

if __name__ == "__main__":
    if not GROQ_API_KEY:
        print("✗ ERREUR: API_KEY n'est pas configurée dans le fichier .env")
        print("Obtenez votre clé API Groq sur: https://console.groq.com/keys")
        exit()
    
    print("🚀 (SCRIPT 1: EXTRACTION avec Groq) 🚀")
    extractor = StagingExtractor(
        api_key=GROQ_API_KEY, 
        reports_folder=REPORTS_FOLDER, 
        db_config=DB_CONFIG
    )
    extractor.process_all_reports(max_reports=None)  
    extractor.close()
    print("\n✅ Extraction terminée. Lancez 'process_staging_to_production.py'.")